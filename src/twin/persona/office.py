"""Office files the twin hands over from work mode: a Word document from light Markdown, a workbook from rows.

Only what a written deliverable needs is understood: ``#`` to ``####`` headings, paragraphs, ``-`` / ``*`` and
``1.`` lists, ``**bold**`` runs and pipe tables. Anything else stays as plain text, so nothing the model writes is
lost. Both writers return the bytes; where they are stored is the caller's business.
"""

from __future__ import annotations

import io
import re
from collections.abc import Sequence

from docx import Document
from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

Cell = str | int | float | bool | None

_HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)、]\s+(.*)$")
_TABLE_ROW = re.compile(r"^\s*\|(.*)\|\s*$")
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
# Excel caps a sheet name at 31 characters and forbids these.
_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
MAX_SHEETS = 20
MAX_ROWS = 10_000
MAX_COLUMNS = 100
COLUMN_WIDTH_CAP = 60


def _runs(paragraph: Paragraph, text: str) -> None:
    """Add ``text`` to ``paragraph``, with ``**…**`` spans in bold."""
    end = 0
    for match in _BOLD.finditer(text):
        if match.start() > end:
            paragraph.add_run(text[end : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        end = match.end()
    if end < len(text):
        paragraph.add_run(text[end:])


def _cells(line: str) -> list[str]:
    row = _TABLE_ROW.match(line)
    assert row is not None
    return [cell.strip() for cell in row.group(1).split("|")]


def _table(document: DocxDocument, rows: list[list[str]]) -> None:
    width = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            cell.text = ""
            _runs(cell.paragraphs[0], row[c] if c < len(row) else "")
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def docx_bytes(title: str, markdown: str) -> bytes:
    """A Word document titled ``title`` (empty: no title) with ``markdown`` as its body."""
    document = Document()
    if title.strip():
        document.add_heading(title.strip(), level=0)
    lines = markdown.replace("\r\n", "\n").split("\n")
    n = 0
    while n < len(lines):
        line = lines[n].rstrip()
        if not line.strip():
            n += 1
            continue
        if _TABLE_ROW.match(line):
            rows: list[list[str]] = []
            while n < len(lines) and (_TABLE_ROW.match(lines[n]) or _TABLE_RULE.match(lines[n])):
                if not _TABLE_RULE.match(lines[n]):
                    rows.append(_cells(lines[n]))
                n += 1
            if rows:
                _table(document, rows)
            continue
        if heading := _HEADING.match(line):
            document.add_heading(heading.group(2).strip(), level=len(heading.group(1)))
        elif bullet := _BULLET.match(line):
            _runs(document.add_paragraph(style="List Bullet"), bullet.group(1))
        elif numbered := _NUMBERED.match(line):
            _runs(document.add_paragraph(style="List Number"), numbered.group(1))
        else:
            _runs(document.add_paragraph(), line.strip())
        n += 1
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _sheet_name(name: str, taken: set[str]) -> str:
    base = _SHEET_FORBIDDEN.sub(" ", name).strip()[:31] or "Sheet"
    candidate, k = base, 2
    while candidate.lower() in taken:
        suffix = f" ({k})"
        candidate, k = base[: 31 - len(suffix)] + suffix, k + 1
    taken.add(candidate.lower())
    return candidate


def xlsx_bytes(sheets: Sequence[tuple[str, Sequence[Sequence[Cell]]]]) -> bytes:
    """A workbook with one sheet per ``(name, rows)``; the first row of each is the bold header. Raises
    ``ValueError`` when there is nothing to write or a sheet is too large."""
    if not sheets:
        raise ValueError("至少要有一个工作表")
    if len(sheets) > MAX_SHEETS:
        raise ValueError(f"工作表不能超过 {MAX_SHEETS} 个")
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    taken: set[str] = set()
    for name, rows in sheets:
        if len(rows) > MAX_ROWS or any(len(r) > MAX_COLUMNS for r in rows):
            raise ValueError(f"每个工作表最多 {MAX_ROWS} 行、{MAX_COLUMNS} 列")
        sheet = workbook.create_sheet(_sheet_name(name, taken))
        widths: dict[int, int] = {}
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row, 1):
                cell = sheet.cell(row=r, column=c, value=value)
                if r == 1:
                    cell.font = Font(bold=True)
                widths[c] = max(widths.get(c, 0), len(str(value)) if value is not None else 0)
        for c, width in widths.items():
            sheet.column_dimensions[get_column_letter(c)].width = min(COLUMN_WIDTH_CAP, max(8, width * 2 + 2))
        if rows:
            sheet.freeze_panes = "A2"
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
