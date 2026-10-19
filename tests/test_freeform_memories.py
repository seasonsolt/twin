from __future__ import annotations

import datetime as dt
import io
import re
import threading
import zipfile
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pydantic import BaseModel
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from typer.testing import CliRunner

from twin.cli import app as cli
from twin.config import Settings
from twin.embed import HashingEmbedder
from twin.llm import FakeLLM
from twin.persona import text as text_module
from twin.persona.chat import PersonaChat
from twin.persona.schema import ChatTurn, SourceKind
from twin.persona.sources import detect_kind, extract_text, parse_chat, parse_upload, pseudonym
from twin.persona.store import PersonaStore
from twin.web import create_app
from twin.web.jobs import JobManager, PersonaProcessing


class Handle:
    def __init__(self, fn: Callable[[], None]) -> None:
        self.fn = fn
        self.cancelled = False

    def cancel(self) -> None:
        self.cancelled = True


class Clock:
    def __init__(self) -> None:
        self.handles: list[Handle] = []
        self.delays: list[float] = []

    def schedule(self, delay: float, fn: Callable[[], None]) -> Handle:
        handle = Handle(fn)
        self.delays.append(delay)
        self.handles.append(handle)
        return handle

    def fire(self) -> None:
        handles, self.handles = self.handles, []
        for handle in handles:
            if not handle.cancelled:
                handle.fn()


def pdf(text: bool) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 10 100 Td (Hello memory) Tj ET")
        page[NameObject("/Contents")] = stream
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_file_text_formats(monkeypatch: pytest.MonkeyPatch) -> None:
    assert extract_text("tiny.PDF", pdf(True)) == "Hello memory"
    document = Document()
    document.add_paragraph("我喜欢读书。")
    row = document.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text, row.cells[1].text = "姓名", "虚构人物"
    buffer = io.BytesIO()
    document.save(buffer)
    assert extract_text("a.docx", buffer.getvalue()) == "我喜欢读书。\n\n姓名\t虚构人物"
    html = "<p>第一段 &amp; 内容</p><script>secret()</script><style>hidden</style><p>第二段<br>换行</p>"
    rendered = extract_text("a.html", html.encode())
    assert "secret" not in rendered and "hidden" not in rendered
    assert "第一段 & 内容" in rendered and "\n\n" in rendered and "第二段" in rendered
    for suffix in (".txt", ".md", ".csv", ".json", ".srt", ".vtt", ".htm"):
        assert extract_text("a" + suffix, "我喜欢咖啡，也喜欢散步。".encode("gbk")) == "我喜欢咖啡，也喜欢散步。"
    assert extract_text("a.txt", b"") == ""
    with pytest.raises(ValueError, match=re.escape("这个 PDF 没有可提取的文字（可能是扫描件）")):
        extract_text("scan.pdf", pdf(False))
    with pytest.raises(ValueError, match="不支持"):
        extract_text("a.exe", b"hello")
    monkeypatch.setattr(text_module, "MAX_FILE_BYTES", 8)
    with pytest.raises(ValueError, match="1 GB"):
        extract_text("big.txt", b"x" * 9)
    monkeypatch.undo()
    with pytest.raises(ValueError, match="无法读取"):
        extract_text("bad.docx", b"broken")


@pytest.mark.parametrize(
    ("name", "text", "kind"),
    [
        ("a.md", "1. 偏好？ 开放 · 2.1\n回答：先看结果", SourceKind.QUESTIONNAIRE),
        ("a.txt", "2025-01-01 10:00 小明：你好\n2025-01-01 10:01 林沐：好", SourceKind.CHAT),
        ("a.csv", "time,sender,content\n2025-01-01,林沐,好", SourceKind.CHAT),
        ("a.json", '{"messages":[{"时间":"2025-01-01","发送人":"林沐","内容":"好"}]}', SourceKind.CHAT),
        ("a.csv", "sender,content\n林沐,好", SourceKind.DOCUMENT),
        ("a.json", '{"text":"文档"}', SourceKind.DOCUMENT),
        ("a.txt", "访谈者：你喜欢什么？\n阿沐：读书", SourceKind.INTERVIEW),
        ("meeting.srt", "1\n00:00:01 --> 00:00:02\n小明：讨论议题", SourceKind.DOCUMENT),
        ("biography.md", "林沐生于某地，他喜欢散步。", SourceKind.DOCUMENT),
        (
            "meeting.pdf",
            "发⾔⼈ 1  00:00:00\n这个方案\n狮⼦  00:00:02\n我来决定。\n发⾔⼈ 2  00:00:04\n好。",
            SourceKind.INTERVIEW,
        ),
        (
            "not-mine.pdf",
            "发⾔⼈ 1  00:00:00\n这个方案\n发⾔⼈ 2  00:00:02\n好。\n发⾔⼈ 1  00:00:04\n确定。",
            SourceKind.INTERVIEW,
        ),
        ("t.txt", "林沐(00:00:03)：先做\n小明(00:00:20)：延后\n发言人1(00:01:02)：同意", SourceKind.INTERVIEW),
        ("t.txt", "[00:00:03] 林沐：先做\n[00:00:20] 小明：延后\n[00:01:02] 发言人1：同意", SourceKind.INTERVIEW),
        (
            "t.txt",
            "小明（00:00:03）\n先做\n小张（00:00:20）\n延后\n小明（00:01:02）\n好\n小张（00:01:09）\n行",
            SourceKind.INTERVIEW,
        ),
        ("t.txt", "林沐：大家好\n第二行\n第三行\n小明：延后\n第二行\n第三行", SourceKind.INTERVIEW),
        ("t.txt", "周会\n小明：先做\n补充\n小张：延后\n小明：好\n小张：行", SourceKind.INTERVIEW),
        # The person's own texts that only look timed or labelled stay documents.
        ("议程.md", "周会议程\n\n签到 09:00\n开场 09:30\n分享 10:00\n午餐 12:00", SourceKind.DOCUMENT),
        ("日记.md", "早上起床 07:30\n想清楚了排期。\n晚饭 19:00\n聊了很多。\n睡觉 23:00", SourceKind.DOCUMENT),
        ("周记.md", "我的周记\n今天整理了思路。\n备注：明天继续\n另外想到一点。\n备注：别忘了", SourceKind.DOCUMENT),
        ("a.md", "我的经历\n\n我喜欢读书。", SourceKind.DOCUMENT),
        ("a.txt", "小明：讨论\n小张：同意", SourceKind.DOCUMENT),
        ("a.txt", "2025-01-01 林沐：好\n续行\n更多续行", SourceKind.DOCUMENT),
        ("a.txt", "", SourceKind.DOCUMENT),
    ],
)
def test_kind_detection(name: str, text: str, kind: SourceKind) -> None:
    assert detect_kind(name, text, Settings(target_name="林沐", target_aliases=["阿沐"])) is kind
    assert kind in set(SourceKind)


def test_upload_dates() -> None:
    settings = Settings(target_name="林沐")
    document = parse_upload("2024-03-02_想法.md", "一些想法".encode(), settings)
    assert document.source.first_date == dt.date(2024, 3, 2)
    today = parse_upload("想法.md", "一些想法".encode(), settings)
    assert today.source.first_date == dt.date.today()
    chat = parse_upload(
        "chat.csv",
        b"time,sender,content\n2024-03-01,\xe6\x9e\x97\xe6\xb2\x90,hi\n,\xe6\x9e\x97\xe6\xb2\x90,hello",
        settings,
    )
    assert all(e.date == dt.date(2024, 3, 1) for e in chat.expressions)


def test_pdf_transcripts_keep_speakers_and_wait_when_the_target_is_not_found(tmp_path: Path) -> None:
    settings = Settings(target_name="林沐", target_aliases=["阿沐"], db_path=tmp_path / "twin.db")
    raw = "发⾔⼈ 1  00:00:00\n怎么看？\n林沐  00:00:02\n先做。\n发⾔⼈ 2  00:00:04\n好。"
    parsed = parse_upload("meeting.txt", raw.encode(), settings)
    assert parsed.source.kind is SourceKind.INTERVIEW
    assert [(e.speaker, e.is_target, e.text) for e in parsed.expressions] == [
        ("发言人 1", False, "怎么看？"),
        ("林沐", True, "先做。"),
        ("发言人 2", False, "好。"),
    ]
    assert parsed.expressions[1].context == "发言人 1：怎么看？"
    without_target = "发⾔⼈ 1  00:00:00\n怎么看？\n发⾔⼈ 2  00:00:02\n先做。\n发⾔⼈ 1  00:00:04\n确定。"
    waiting = parse_upload("not-mine.txt", without_target.encode(), settings)
    assert waiting.source.kind is SourceKind.INTERVIEW and waiting.source.media_status == "needs_name"
    assert waiting.source.n_target == 0 and not any(e.is_target or e.context for e in waiting.expressions)


@pytest.mark.parametrize(
    "raw",
    [
        "林沐(00:00:03)：先做。\n小明(00:00:20)：延后？\n林沐(00:01:02)：不延。",
        "[00:00:03] 林沐：先做。\n[00:00:20] 小明：延后？\n[00:01:02] 林沐：不延。",
        "林沐 00:00:03\n先做。\n小明 00:00:20\n延后？\n林沐 00:01:02\n不延。",
    ],
)
def test_timed_transcript_turns_are_attributed_by_speaker(raw: str, tmp_path: Path) -> None:
    parsed = parse_upload("会议.txt", raw.encode(), Settings(target_name="林沐", db_path=tmp_path / "twin.db"))
    assert parsed.source.kind is SourceKind.INTERVIEW
    assert [(e.speaker, e.is_target, e.text) for e in parsed.expressions] == [
        ("林沐", True, "先做。"),
        ("小明", False, "延后？"),
        ("林沐", True, "不延。"),
    ]


def test_timed_transcripts_keep_colon_lines_inside_a_turn(tmp_path: Path) -> None:
    raw = "林沐 00:00:03\n注意：先做A\n再做B\n小明 00:00:20\n延后\n林沐 00:01:02\n好"
    parsed = parse_upload("会议.txt", raw.encode(), Settings(target_name="林沐", db_path=tmp_path / "twin.db"))
    assert [(e.speaker, e.text) for e in parsed.expressions] == [
        ("林沐", "注意：先做A\n再做B"),
        ("小明", "延后"),
        ("林沐", "好"),
    ]


def test_chat_lines_with_a_clock_time_continue_the_message(tmp_path: Path) -> None:
    raw = "2026-09-01 10:00 林沐：明天几点开会？\n明天 10:00\n小明 2026-09-01 10:05\n好的"
    parsed = parse_chat("群.txt", raw, Settings(target_name="林沐", db_path=tmp_path / "twin.db"))
    assert [(e.speaker, e.text) for e in parsed.expressions] == [
        ("林沐", "明天几点开会？\n明天 10:00"),
        ("小明", "好的"),
    ]


def test_pdf_text_folds_compatibility_ideographs_but_keeps_full_width_punctuation(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class Page:
        def extract_text(self) -> str:
            return "狮⼦：我⼦⼥都在，好！"

    class Reader:
        def __init__(self, _: Any) -> None:
            self.pages = [Page()]

    monkeypatch.setattr("pypdf.PdfReader", Reader)
    assert extract_text("m.pdf", b"%PDF") == "狮子：我子女都在，好！"


def test_debounce_and_exactly_one_follow_up() -> None:
    clock = Clock()
    jobs = JobManager(1, lambda e: str(e))
    started, release, finished = threading.Event(), threading.Event(), threading.Event()
    calls: list[int] = []

    def build(log: Callable[[str], None]) -> dict[str, int]:
        calls.append(1)
        started.set()
        assert release.wait(5)
        return {"items": 0}

    processing = PersonaProcessing(jobs, build, lambda _: finished.set(), schedule=clock.schedule)
    assert processing.view()["state"] == "idle"
    for _ in range(5):
        processing.queue()
    assert processing.view()["state"] == "queued"
    assert all(delay == 3 for delay in clock.delays)
    assert sum(not handle.cancelled for handle in clock.handles) == 1
    clock.fire()
    assert started.wait(5)
    assert processing.view()["state"] == "running"
    for _ in range(10):
        processing.queue()
    release.set()
    assert finished.wait(5)
    assert processing.view()["state"] == "queued"
    assert sum(not handle.cancelled for handle in clock.handles) == 1
    finished.clear()
    clock.fire()
    assert finished.wait(5)
    assert processing.view()["state"] == "idle"
    assert len(calls) == 2
    assert processing.view()["last_finished_at"] and processing.view()["last_error"] is None
    processing.close()


def test_processing_adopts_external_build_and_recovers_failure() -> None:
    clock, jobs = Clock(), JobManager(1, lambda e: str(e))
    release, finished = threading.Event(), threading.Event()
    external = jobs.submit("persona_build", "external", lambda _: release.wait(5))
    processing = PersonaProcessing(
        jobs,
        lambda _: (_ for _ in ()).throw(ValueError("safe failure")),
        lambda _: finished.set(),
        schedule=clock.schedule,
    )
    processing.queue()
    assert processing.start() is external
    assert processing.view()["job_id"] == external.job_id
    release.set()
    assert finished.wait(5)
    assert processing.view()["state"] == "queued"
    finished.clear()
    clock.fire()
    assert finished.wait(5)
    assert processing.view()["state"] == "idle"
    assert "safe failure" in processing.view()["last_error"]
    processing.close()


def handler(system: str, user: str, schema: type[BaseModel]) -> dict[str, Any]:
    if schema.__name__ == "ExtractDraft":
        if "钱进了账户" not in user:
            return {"items": []}
        n = int(re.findall(r"\[(\d+)\]", user)[0])
        return {"items": [{"facet_id": "2.1", "statement": "看重回款", "quotes": [{"n": n, "quote": "钱进了账户"}]}]}
    if schema.__name__ == "MergeDraft":
        return {"items": [{"statement": "看重回款", "candidate_ids": re.findall(r"\[(pc_[0-9a-f]+)\]", user)}]}
    raise AssertionError(schema)


@pytest.fixture
def web(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[tuple[TestClient, Clock, Settings]]:
    clock = Clock()
    monkeypatch.setattr(PersonaProcessing, "_timer", staticmethod(clock.schedule))
    settings = Settings(target_name="林沐", db_path=tmp_path / "persona.db")
    app = create_app(settings, llm_factory=lambda: FakeLLM(handler), embedder_factory=HashingEmbedder)
    with TestClient(app, base_url="http://127.0.0.1", headers={"X-Twin": "1"}) as client:
        yield client, clock, settings


def finish_build(client: TestClient) -> None:
    response = client.post("/api/persona/build")
    assert response.status_code == 202
    # The manager exposes a wait method; the HTTP contract remains polled by the browser.
    from time import monotonic, sleep

    deadline = monotonic() + 5
    while monotonic() < deadline:
        if client.get("/api/persona/processing").json()["state"] == "idle":
            return
        sleep(0.001)
    raise AssertionError("build did not finish")


def test_notes_import_status_privacy_delete_and_chat(web: tuple[TestClient, Clock, Settings]) -> None:
    client, _clock, _settings = web
    assert client.get("/api/persona/processing").json()["state"] == "idle"
    response = client.post("/api/persona/chat", json={"messages": [{"role": "user", "content": "你好"}]})
    assert response.status_code == 200 and response.json()["abstain"]
    assert "还没有添加记忆" in response.json()["reply"]
    note = client.post("/api/persona/notes", json={"text": "一些普通文字"}).json()
    assert note["title"].startswith("笔记 ")
    assert note["detected_kind_label"] == "笔记" and note["first_date"] == dt.date.today().isoformat()
    assert client.post("/api/persona/notes", json={"text": ""}).status_code == 400
    assert client.post("/api/persona/notes", json={"text": " "}).status_code == 400
    assert client.post("/api/persona/notes", json={"text": "x" * 20001}).status_code == 400
    assert client.get("/api/persona/processing").json()["state"] == "queued"
    assert client.get("/api/persona/sources").json()["items"][0]["status"] == "processing"
    reply = client.post("/api/persona/chat", json={"messages": [{"role": "user", "content": "你好"}]})
    assert reply.status_code == 200 and "处理中" in reply.json()["reply"]
    finish_build(client)
    rows = client.get("/api/persona/sources").json()["items"]
    assert rows[0]["status"] == "nothing_found" and rows[0]["remembered"] == 0 and "preview" not in rows[0]
    uploads = [
        (
            "files",
            (
                "folder/2024-01-02_聊天.txt",
                "2024-01-02 10:00 王五：你好\n2024-01-02 10:01 林沐：王五说钱进了账户".encode(),
            ),
        ),
        ("files", ("folder/.hidden.txt", b"hidden")),
        ("files", (".private/a.txt", b"hidden")),
        ("files", ("folder/a.exe", b"unsupported")),
    ]
    result = client.post("/api/persona/import", files=uploads).json()
    assert len(result["imported"]) == 1 and len(result["skipped"]) == 3
    source = result["imported"][0]
    assert source["origin"] == "2024-01-02_聊天.txt" and source["detected_kind"] == "chat"
    assert source["detected_kind_label"] == "聊天记录"
    preview = client.get(f"/api/persona/sources/{source['source_id']}/text")
    assert "王五" not in preview.text and pseudonym("王五") in preview.text
    assert client.get("/api/persona/sources/missing/text").status_code == 404
    finish_build(client)
    remembered = next(
        s for s in client.get("/api/persona/sources").json()["items"] if s["source_id"] == source["source_id"]
    )
    assert remembered["status"] == "remembered" and remembered["remembered"] == 1
    assert client.get("/api/persona/processing").json()["last_finished_at"]
    assert client.delete(f"/api/persona/sources/{source['source_id']}").status_code == 200
    assert client.get("/api/persona/processing").json()["state"] == "queued"
    assert client.delete("/api/persona/sources/missing").status_code == 404
    finish_build(client)
    assert client.get("/api/persona/items").json()["items"] == []
    title = client.post("/api/persona/notes", json={"text": "x" * 20000, "title": "我的想法"}).json()
    assert title["title"] == "我的想法"
    assert len(client.get(f"/api/persona/sources/{title['source_id']}/text").text) == 20000
    finish_build(client)


def test_import_limits(web: tuple[TestClient, Clock, Settings], monkeypatch: pytest.MonkeyPatch) -> None:
    from twin.web import persona

    client, _clock, _settings = web
    assert "1 GB" in persona._too_large().detail
    monkeypatch.setattr(persona, "MAX_UPLOAD_BYTES", 1000)
    monkeypatch.setattr(persona, "MAX_UPLOAD_FILES", 2)
    big = client.post("/api/persona/import", files={"files": ("a.txt", b"x" * 2000)})
    assert big.status_code == 413
    body = b'--b\r\nContent-Disposition: form-data; name="files"; filename="a.txt"\r\n\r\n' + b"x" * 2000
    chunked = client.post(
        "/api/persona/import",
        content=iter([body[:500], body[500:]]),
        headers={"content-type": "multipart/form-data; boundary=b"},
    )
    assert chunked.status_code == 413 and "content-length" not in chunked.request.headers
    many = client.post("/api/persona/import", files=[("files", (f"{i}.txt", b"hi")) for i in range(3)])
    assert many.status_code == 400 and "一次最多上传 2 个文件" in many.json()["detail"]
    none = client.post("/api/persona/import", data={"other": "x"}, files={"note": ("a.txt", b"hi")})
    assert none.status_code == 400 and "没有收到文件" in none.json()["detail"]
    broken = client.post("/api/persona/import", content=b"x", headers={"content-type": "multipart/form-data"})
    assert broken.status_code == 400 and "无法解析上传内容" in broken.json()["detail"]


def test_processing_endpoint_running_follow_up_and_startup_resume(
    web: tuple[TestClient, Clock, Settings], monkeypatch: pytest.MonkeyPatch
) -> None:
    from time import monotonic, sleep

    from twin.web.persona import run_persona_build

    client, clock, settings = web
    started, release = threading.Event(), threading.Event()
    calls: list[int] = []

    def block(*args: Any) -> dict[str, Any]:
        calls.append(1)
        args[-1]("[1/2] 抽取并合并人格档案")
        args[-1]("ok persona extract ch_test")
        started.set()
        assert release.wait(5)
        return run_persona_build(*args)

    monkeypatch.setattr("twin.web.persona.run_persona_build", block)
    client.post("/api/persona/notes", json={"text": "first"})
    clock.fire()
    assert started.wait(5)
    running = client.get("/api/persona/processing").json()
    assert running["state"] == "running" and running["job_id"]
    assert running["stage"] == {"current": 1, "total": 2, "label": "抽取并合并人格档案"}
    assert running["tally"] == {"done": 1, "failed": 0, "total": None}
    # what the page needs to tick between polls: the start, the server's clock and the live activity
    # each carries its offset, so a viewer in any time zone counts the same seconds
    started, now = (dt.datetime.fromisoformat(running[key]) for key in ("started", "now"))
    assert started.tzinfo is not None and now.tzinfo is not None and now >= started
    assert running["activity"]["waiting"] == 0 and running["activity"]["retries"] == 0
    assert dt.datetime.fromisoformat(running["activity"]["last_progress_at"]) >= started
    for number in range(5):
        client.post("/api/persona/notes", json={"text": f"next {number}"})
    release.set()
    deadline = monotonic() + 5
    while client.get("/api/persona/processing").json()["state"] != "queued":
        assert monotonic() < deadline
        sleep(0.001)
    clock.fire()
    while client.get("/api/persona/processing").json()["state"] != "idle":
        assert monotonic() < deadline
        sleep(0.001)
    assert len(calls) == 2
    assert all(s["status"] == "nothing_found" for s in client.get("/api/persona/sources").json()["items"])
    client.post("/api/persona/notes", json={"text": "saved before shutdown"})
    second_app = create_app(settings, llm_factory=lambda: FakeLLM(handler), embedder_factory=HashingEmbedder)
    with TestClient(second_app, base_url="http://127.0.0.1") as restarted:
        assert restarted.get("/api/persona/processing").json()["state"] == "queued"


def test_failed_build_status_and_private_error(
    web: tuple[TestClient, Clock, Settings], monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    client, _, settings = web
    secret = "NEVER-LOG-PERSONAL-TEXT"

    def fail(*args: Any) -> Any:
        raise RuntimeError(secret)

    monkeypatch.setattr("twin.web.persona.run_persona_build", fail)
    source = client.post("/api/persona/notes", json={"text": secret}).json()
    finish_build(client)
    assert client.get("/api/persona/sources").json()["items"][0]["status"] == "failed"
    assert client.get("/api/persona/processing").json()["last_error"]
    assert secret not in caplog.text
    with PersonaStore(settings.db_path) as store:
        store.processing_result([source["source_id"]], "old-version", "old error")
        assert store.get_meta(f"source_error:{source['source_id']}") != "old error"
    result = client.post("/api/persona/import", files={"files": ("scan.pdf", pdf(False))}).json()
    assert result["imported"] == [] and "可能是扫描件" in result["skipped"][0]["reason"]


def test_no_profile_chat_does_not_call_models() -> None:
    settings = Settings()
    llm = FakeLLM(lambda *_: (_ for _ in ()).throw(AssertionError("must not call")))
    with PersonaStore(":memory:") as store:
        reply = PersonaChat(store, llm, HashingEmbedder(), settings).reply([ChatTurn(role="user", content="你好")])
        assert reply.abstain and reply.mode == "abstain" and not llm.calls


def test_cli_note_and_auto_kind(tmp_path: Path) -> None:
    config = tmp_path / "twin.toml"
    config.write_text('target_name = "林沐"\ndb_path = "persona.db"\n')
    path = tmp_path / "资料.txt"
    path.write_bytes("2024-01-01 10:00 林沐：我喜欢读书".encode("gbk"))
    runner = CliRunner()
    result = runner.invoke(cli, ["--config", str(config), "persona", "import", str(path)])
    assert result.exit_code == 0 and "聊天记录" in result.output
    note = runner.invoke(cli, ["--config", str(config), "persona", "note", "我喜欢散步"])
    assert note.exit_code == 0 and "已添加 笔记" in note.output
    with PersonaStore(tmp_path / "persona.db") as store:
        assert len(store.list_sources()) == 2


def epub(chapters: dict[str, str], spine: list[str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as book:
        book.writestr("mimetype", "application/epub+zip")
        book.writestr(
            "META-INF/container.xml",
            '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
            '<rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>",
        )
        items = "".join(
            f'<item id="{name}" href="Text/{name}.xhtml" media-type="application/xhtml+xml"/>' for name in chapters
        )
        refs = "".join(f'<itemref idref="{name}"/>' for name in spine)
        book.writestr(
            "OEBPS/content.opf",
            f'<package xmlns="http://www.idpf.org/2007/opf" version="3.0"><manifest>{items}'
            '<item id="css" href="style.css" media-type="text/css"/></manifest>'
            f'<spine>{refs}<itemref idref="css"/></spine></package>',
        )
        book.writestr("OEBPS/style.css", "p { color: red }")
        for name, body in chapters.items():
            book.writestr(
                f"OEBPS/Text/{name}.xhtml", f"<html><head><style>p{{}}</style></head><body>{body}</body></html>"
            )
    return buffer.getvalue()


def test_epub_text_follows_the_reading_order() -> None:
    data = epub({"b": "<h1>第二章</h1><p>后来我搬到了海边。</p>", "a": "<p>我在山里长大。</p>"}, ["a", "b"])
    text = extract_text("回忆录.EPUB", data)
    assert text.index("我在山里长大") < text.index("第二章") < text.index("后来我搬到了海边")
    assert "color" not in text and "p{}" not in text
    with pytest.raises(ValueError, match="无法读取电子书"):
        extract_text("bad.epub", b"broken")
    with pytest.raises(ValueError, match="没有可提取的文字"):
        extract_text("empty.epub", epub({"a": "<img src='x.png'/>"}, ["a"]))
